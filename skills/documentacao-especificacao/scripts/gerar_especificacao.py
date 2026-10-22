#!/usr/bin/env python3
"""Gera o Documento de especificação ADS a partir do modelo .docx e de um JSON de conteúdo.

Uso:
    python3 gerar_especificacao.py conteudo.json [--modelo caminho/modelo.docx]

Mantém do modelo: capa (logotipo e faixa), estilos, cabeçalho/rodapé, sumário (campo TOC) e a
tabela "Controle de Versões" (preenchida com `versoes`). Todo o corpo a partir do primeiro
Título 1 após o sumário é substituído pelos `blocos` do JSON, com a formatação do modelo
(Calibri, cor 414141, tabelas com cabeçalho cinza D9D9D9, marcadores do modelo).

Capa (opcional): {"capa": {"projeto": "...", "time_ti": "Arquitetura e desenvolvimento de software",
  "analista": "...", "data_publicacao": "outubro de 2026", "data_iso": "2026-10-19",
  "departamento": "...", "solicitante": "..."}}

Blocos aceitos (campo "t"):
  h1, h2, h3, h4        {"t":"h2","texto":"Consulta de solicitações"}
  p                     {"t":"p","texto":"Texto com **negrito** inline"}
  rotulo                {"t":"rotulo","texto":"Cabeçalho"}            -> parágrafo em negrito
  li                    {"t":"li","texto":"**Campo:** descrição","nivel":0}
  tabela                {"t":"tabela","cabecalho":["A","B"],"linhas":[["1","2"]],"larguras_cm":[4,12]}
  chave_valor           {"t":"chave_valor","linhas":[["URL","https://..."],["Método","POST"]]}
  img                   {"t":"img","caminho":"03-capturas/TL-001.png","largura_cm":15,"legenda":"TL-001"}
  codigo                {"t":"codigo","texto":"{\n  \"Name\": \"X\"\n}"}
  cenario               {"t":"cenario","titulo":"CT-001 — Verificar se ...","passos":["DADO QUE ...","E ...","QUANDO ...","ENTÃO ..."]}
  quebra                {"t":"quebra"}  -> quebra de página
"""
import argparse
import copy
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
MODELO_PADRAO = os.path.join(AQUI, "..", "assets", "modelo-especificacao-ads.docx")
COR = RGBColor(0x41, 0x41, 0x41)
FONTE = "Calibri"
NUM_MARCADOR = "1"  # numId dos marcadores usados nos requisitos do modelo


def fmt_run(run, negrito=False, italico=False, tamanho=None, mono=False):
    run.font.name = "Consolas" if mono else FONTE
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    nome = "Consolas" if mono else FONTE
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), nome)
    run.font.color.rgb = COR
    run.bold = negrito or None
    run.italic = italico or None
    if tamanho:
        run.font.size = Pt(tamanho)


def add_texto(par, texto, negrito_base=False, tamanho=None):
    """Aceita **negrito** inline. Marcas [pendente] etc. ficam em itálico."""
    partes = re.split(r"(\*\*[^*]+\*\*|\[(?:pendente|inferido|divergente|sugerido|externo)\])", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            fmt_run(par.add_run(parte[2:-2]), negrito=True, tamanho=tamanho)
        elif re.fullmatch(r"\[(pendente|inferido|divergente|sugerido|externo)\]", parte):
            fmt_run(par.add_run(parte), negrito=True, italico=True, tamanho=tamanho)
        else:
            fmt_run(par.add_run(parte), negrito=negrito_base, tamanho=tamanho)


def set_num(par, nivel):
    ppr = par._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(nivel))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), NUM_MARCADOR)
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)


def sombrear(celula, cor="D9D9D9"):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def texto_celula(celula, texto, negrito=False):
    celula.text = ""
    linhas = str(texto).split("\n")
    for i, linha in enumerate(linhas):
        par = celula.paragraphs[0] if i == 0 else celula.add_paragraph()
        ppr = par._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLine"), "0")
        ppr.append(ind)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "left")
        ppr.append(jc)
        par.paragraph_format.space_after = Pt(0)
        add_texto(par, linha, negrito_base=negrito, tamanho=10)


def tabela(doc, cabecalho, linhas, larguras=None, cabecalho_coluna=False):
    ncol = len(cabecalho) if cabecalho else len(linhas[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = doc.styles["Table Grid"]
    if cabecalho:
        r = t.add_row()
        for c, txt in zip(r.cells, cabecalho):
            texto_celula(c, txt, negrito=True)
            sombrear(c)
    for linha in linhas:
        r = t.add_row()
        for j, (c, txt) in enumerate(zip(r.cells, list(linha) + [""] * (ncol - len(linha)))):
            texto_celula(c, txt, negrito=cabecalho_coluna and j == 0)
            if cabecalho_coluna and j == 0:
                sombrear(c)
    if larguras:
        for r in t.rows:
            for c, w in zip(r.cells, larguras):
                c.width = Cm(w)
    doc.add_paragraph()
    return t


def limpar_corpo(doc):
    """Remove do primeiro Título 1 depois do sumário até o fim (preserva sectPr)."""
    corpo = doc.element.body
    elementos = list(corpo)
    inicio = None
    for i, el in enumerate(elementos):
        if el.tag == qn("w:p"):
            estilo = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
            if estilo is not None and txt.strip() == "Objetivo":
                inicio = i
                break
    if inicio is None:
        sys.exit("Não encontrei o início do corpo (Título 1 'Objetivo') no modelo.")
    for el in elementos[inicio:]:
        if el.tag != qn("w:sectPr"):
            corpo.remove(el)


def preencher_versoes(doc, versoes):
    t = doc.tables[0]
    modelo_linha = copy.deepcopy(t.rows[1]._tr)
    for r in list(t.rows)[1:]:
        t._tbl.remove(r._tr)
    for v in versoes:
        tr = copy.deepcopy(modelo_linha)
        t._tbl.append(tr)
        linha = t.rows[-1]
        valores = [v.get("autor", ""), v.get("versao", ""), v.get("data", ""), v.get("resumo", "")]
        for c, val in zip(linha.cells, valores):
            pars = c.paragraphs
            for p in pars[1:]:
                p._p.getparent().remove(p._p)
            p = pars[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            fmt_run(p.add_run(str(val)), tamanho=10)


CAPA_ALIAS = {
    "projeto": "Projeto",
    "time_ti": "Time de TI",
    "analista": "Analista",
    "data_publicacao": "Data de Publicação",
    "departamento": "Departamento",
    "solicitante": "Solicitante",
}


def _xml_esc(t):
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def preencher_capa(doc, capa):
    """Preenche os controles de conteúdo da capa e os repositórios a que estão vinculados."""
    if not capa:
        return
    import re as _re
    corpo = doc.element.body
    for chave, alias in CAPA_ALIAS.items():
        if chave not in capa:
            continue
        valor = str(capa[chave])
        for sdt in corpo.iter(qn("w:sdt")):
            al = sdt.find(qn("w:sdtPr") + "/" + qn("w:alias"))
            if al is None or al.get(qn("w:val")) != alias:
                continue
            cont = sdt.find(qn("w:sdtContent"))
            ts = list(cont.iter(qn("w:t")))
            if not ts:
                continue
            ts[0].text = valor
            for t in ts[1:]:
                t.text = ""
            ph = sdt.find(qn("w:sdtPr") + "/" + qn("w:showingPlcHdr"))
            if ph is not None:
                ph.getparent().remove(ph)
    cp = doc.core_properties
    if "projeto" in capa:
        cp.title = str(capa["projeto"])
    if "time_ti" in capa:
        cp.subject = str(capa["time_ti"])
    if "analista" in capa:
        cp.author = str(capa["analista"])
    for parte in doc.part.package.iter_parts():
        nome = str(parte.partname)
        if nome == "/docProps/app.xml" and "departamento" in capa:
            xml = parte.blob.decode("utf-8")
            xml = _re.sub(r"<Company>[^<]*</Company>", f"<Company>{_xml_esc(capa['departamento'])}</Company>", xml)
            parte._blob = xml.encode("utf-8")
        elif nome.startswith("/customXml/item") and b"CoverPageProperties" in parte.blob:
            xml = parte.blob.decode("utf-8")
            if "solicitante" in capa:
                xml = _re.sub(r"<CompanyAddress>[^<]*</CompanyAddress>", f"<CompanyAddress>{_xml_esc(capa['solicitante'])}</CompanyAddress>", xml)
            if "data_iso" in capa:
                xml = _re.sub(r"<PublishDate>[^<]*</PublishDate>", f"<PublishDate>{capa['data_iso']}T00:00:00</PublishDate>", xml)
            parte._blob = xml.encode("utf-8")


def reconstruir_sumario(doc, titulos):
    """Substitui as entradas em cache do sumário pelos títulos gerados (sem número de página).

    O campo TOC é mantido; com updateFields=true o Word recalcula páginas ao abrir.
    """
    alvo = None
    for sdt in doc.element.body.iter(qn("w:sdt")):
        g = sdt.find(".//" + qn("w:docPartGallery"))
        if g is not None and "Contents" in (g.get(qn("w:val")) or ""):
            alvo = sdt
            break
    if alvo is None:
        return
    cont = alvo.find(qn("w:sdtContent"))
    filhos = list(cont)
    cabecalho = filhos[0]
    for f in filhos[1:]:
        cont.remove(f)

    def run(xml_interno):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), FONTE)
        cor = OxmlElement("w:color")
        cor.set(qn("w:val"), "414141")
        rpr.append(rf)
        rpr.append(cor)
        r.append(rpr)
        r.append(xml_interno)
        return r

    def fld(tipo):
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), tipo)
        return e

    estilos = {1: "Sumrio1", 2: "Sumrio2", 3: "Sumrio3"}
    for i, (nivel, texto) in enumerate(titulos or [(1, "Atualize o sumário (F9)")]):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), estilos.get(nivel, "Sumrio3"))
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "10456")
        tabs.append(tab)
        ppr.append(ps)
        ppr.append(tabs)
        p.append(ppr)
        if i == 0:
            p.append(run(fld("begin")))
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            p.append(run(it))
            p.append(run(fld("separate")))
        t = OxmlElement("w:t")
        t.text = texto
        p.append(run(t))
        cont.append(p)
    p = OxmlElement("w:p")
    p.append(run(fld("end")))
    cont.append(p)


def atualizar_campos_ao_abrir(doc):
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("conteudo")
    ap.add_argument("--modelo", default=MODELO_PADRAO)
    a = ap.parse_args()
    with open(a.conteudo, encoding="utf-8") as f:
        dados = json.load(f)
    base = os.path.dirname(os.path.abspath(a.conteudo))

    doc = Document(a.modelo)
    preencher_capa(doc, dados.get("capa"))
    preencher_versoes(doc, dados.get("versoes", []))
    limpar_corpo(doc)

    estilos_titulo = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "h4": "Heading 4"}
    titulos = [(int(b["t"][1]), b["texto"]) for b in dados["blocos"] if b["t"] in ("h1", "h2", "h3")]
    reconstruir_sumario(doc, titulos)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # corpo começa em nova página
    for b in dados["blocos"]:
        t = b["t"]
        if t in estilos_titulo:
            doc.add_paragraph(b["texto"], style=estilos_titulo[t])
        elif t == "p":
            add_texto(doc.add_paragraph(), b["texto"])
        elif t == "rotulo":
            add_texto(doc.add_paragraph(), b["texto"], negrito_base=True)
        elif t == "li":
            p = doc.add_paragraph(style="List Paragraph")
            set_num(p, b.get("nivel", 0))
            add_texto(p, b["texto"])
        elif t == "tabela":
            tabela(doc, b.get("cabecalho"), b["linhas"], b.get("larguras_cm"))
        elif t == "chave_valor":
            tabela(doc, None, b["linhas"], b.get("larguras_cm", [3.5, 12.5]), cabecalho_coluna=True)
        elif t == "img":
            caminho = b["caminho"] if os.path.isabs(b["caminho"]) else os.path.join(base, b["caminho"])
            if not os.path.exists(caminho):
                add_texto(doc.add_paragraph(), f"[pendente] captura não encontrada: {b['caminho']}")
                continue
            pimg = doc.add_paragraph()
            pimg.paragraph_format.first_line_indent = Cm(0)
            pimg.paragraph_format.left_indent = Cm(0)
            pimg.add_run().add_picture(caminho, width=Cm(b.get("largura_cm", 15)))
            if b.get("legenda"):
                p = doc.add_paragraph()
                fmt_run(p.add_run(b["legenda"]), italico=True, tamanho=9)
        elif t == "codigo":
            for linha in b["texto"].split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                fmt_run(p.add_run(linha), tamanho=9, mono=True)
            doc.add_paragraph()
        elif t == "cenario":
            p = doc.add_paragraph(style="List Paragraph")
            set_num(p, 0)
            add_texto(p, b["titulo"], negrito_base=True)
            for passo in b["passos"]:
                add_texto(doc.add_paragraph(), passo)
        elif t == "quebra":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            sys.exit(f"Bloco desconhecido: {t}")

    atualizar_campos_ao_abrir(doc)
    saida = dados.get("saida", "especificacao.docx")
    saida = saida if os.path.isabs(saida) else os.path.join(base, saida)
    doc.save(saida)
    print(saida)


if __name__ == "__main__":
    main()
